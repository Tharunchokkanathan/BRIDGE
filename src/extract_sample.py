import os
import csv
import time
import pandas as pd

CANDIDATE_INPUTS = [
    os.path.join("data", "processed", "amazon_delivery_ml_features_v1.parquet"),
    os.path.join("data", "processed", "amazon_delivery_ml_features_v1.csv"),
    os.path.join("..", "FDSE_PREVIOUS_BACKUP", "data", "processed", "amazon_delivery_ml_features_v1.parquet"),
    os.path.join("..", "FDSE_PREVIOUS_BACKUP", "data", "processed", "amazon_delivery_ml_features_v1.csv"),
    "amazon_delivery_ml_features_v1.xlsx",
]
OUTPUT_DIR = "data"
OUTPUT_CSV = os.path.join(OUTPUT_DIR, "amazon_delivery_ml_features_sample.csv")
TARGET_ROWS = 100000

def extract_sample():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    os.makedirs("models", exist_ok=True)
    
    source_path = None
    for candidate in CANDIDATE_INPUTS:
        if os.path.exists(candidate):
            source_path = candidate
            break
            
    if not source_path:
        print(f"[WARNING] No candidate source dataset found in {CANDIDATE_INPUTS}.")
        if os.path.exists(OUTPUT_CSV):
            print(f"[INFO] Existing sample {OUTPUT_CSV} ({os.path.getsize(OUTPUT_CSV)/1024/1024:.2f} MB) is available.")
        return

    print(f"Reading sample from {source_path}...")
    t0 = time.time()
    
    if source_path.endswith(".parquet"):
        df = pd.read_parquet(source_path)
        sample_df = df.head(TARGET_ROWS)
        sample_df.to_csv(OUTPUT_CSV, index=False)
        row_count = len(sample_df)
    elif source_path.endswith(".csv"):
        sample_df = pd.read_csv(source_path, nrows=TARGET_ROWS)
        sample_df.to_csv(OUTPUT_CSV, index=False)
        row_count = len(sample_df)
    elif source_path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(source_path, read_only=True, data_only=True)
        sheet = wb.active
        row_count = 0
        with open(OUTPUT_CSV, "w", newline="", encoding="utf-8") as f_out:
            writer = None
            for row in sheet.iter_rows(values_only=True):
                if writer is None:
                    writer = csv.writer(f_out)
                    writer.writerow(row)
                    continue
                writer.writerow(row)
                row_count += 1
                if row_count >= TARGET_ROWS:
                    break
        wb.close()
        
    total_time = time.time() - t0
    file_size_mb = os.path.getsize(OUTPUT_CSV) / (1024 * 1024)
    print(f"[SUCCESS] Extracted {row_count:,} rows into {OUTPUT_CSV} ({file_size_mb:.2f} MB) in {total_time:.2f}s.")

if __name__ == "__main__":
    extract_sample()
